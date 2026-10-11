"""Microsoft Excel connector supporting workbooks (.xlsx, .xls), sheets, folders, and remote endpoints."""

from __future__ import annotations

import io
from typing import Any

import polars as pl

from det.connectors.contract import (
    BaseConnector,
    ConfigField,
    ConnectionTestResult,
    ConnectorGroup,
    DiscoveredObject,
    PreviewResult,
    SchemaResult,
)
from det.connectors.errors import DiscoveryError, PreviewError, SchemaInferenceError
from det.connectors.files.file_utils import (
    discover_files_endpoint,
    fetch_remote_bytes,
    format_preview_rows,
    map_polars_dtype_to_metadata,
    resolve_local_file_paths,
    resolve_protocol,
    test_connection_endpoint,
)


class ExcelConnector(BaseConnector):
    id = "excel"
    name = "Excel"
    group = ConnectorGroup.FILES
    description = "Microsoft Excel workbooks (.xlsx, .xls), multi-sheet datasets, and workbook folders."

    source_fields = [
        ConfigField(
            name="protocol",
            label="Protocol / Storage Type",
            type="select",
            default="local",
            options=[
                {"label": "Local Filesystem", "value": "local"},
                {"label": "Remote HTTP / HTTPS", "value": "http"},
                {"label": "Remote FTP Server", "value": "ftp"},
                {"label": "Remote SFTP Server (SSH)", "value": "sftp"},
            ],
            description="Storage protocol for Excel workbook files.",
        ),
        ConfigField(
            name="base_path",
            label="Base Directory or URL",
            type="folder",
            placeholder="/reports or https://reports.example.com",
            description="Root folder or base URL containing Excel workbooks.",
        ),
        ConfigField(
            name="host",
            label="Host",
            type="string",
            placeholder="ftp.example.com",
            description="Host for remote FTP/SFTP connections.",
            advanced=True,
        ),
        ConfigField(
            name="port",
            label="Port",
            type="integer",
            placeholder="21 or 22",
            description="Port number for FTP (21) or SFTP (22).",
            advanced=True,
        ),
        ConfigField(
            name="username",
            label="Username",
            type="string",
            placeholder="reports_user",
            description="Authentication username.",
            advanced=True,
        ),
        ConfigField(
            name="password",
            label="Password / Secret",
            type="password",
            secret=True,
            description="Password or secret token for remote access.",
            advanced=True,
        ),
    ]

    table_fields = [
        ConfigField(
            name="file_selection",
            label="File Selection Mode",
            type="select",
            default="auto",
            options=[
                {"label": "Auto-detect (File or Folder)", "value": "auto"},
                {"label": "Single Workbook File", "value": "single"},
                {"label": "Multiple Files in Folder", "value": "folder"},
                {"label": "Nested Subfolders (Recursive)", "value": "recursive"},
            ],
            description="Read a single workbook or combine multiple workbooks from a folder.",
        ),
        ConfigField(
            name="path",
            label="Workbook Path or Directory",
            type="string",
            required=True,
            placeholder="quarterly_report.xlsx or financial_models",
            description="Path to Excel workbook (.xlsx, .xls), directory, or remote URL.",
        ),
        ConfigField(
            name="file_pattern",
            label="File Filter Pattern",
            type="string",
            default="*.xlsx",
            placeholder="*.xlsx",
            description="Glob pattern when scanning folders of workbooks.",
            advanced=True,
        ),
        ConfigField(
            name="sheet_name",
            label="Sheet Name",
            type="string",
            placeholder="Sheet1 (defaults to first worksheet)",
            description="Specific worksheet name to read. Defaults to the first sheet.",
            advanced=False,
        ),
        ConfigField(
            name="has_header",
            label="First Row as Header",
            type="boolean",
            default=True,
            description="Whether the first non-skipped row contains column header names.",
            advanced=False,
        ),
        ConfigField(
            name="skip_rows",
            label="Skip Leading Rows",
            type="integer",
            default=0,
            placeholder="0",
            description="Number of header/metadata rows to skip before column headers.",
            advanced=True,
        ),
        ConfigField(
            name="columns",
            label="Columns Projection",
            type="string",
            placeholder="col_a, col_b",
            description="Optional comma-separated list of column names to project.",
            advanced=True,
        ),
    ]

    def test_connection(
        self,
        source_config: dict[str, Any],
        credentials: dict[str, Any],
    ) -> ConnectionTestResult:
        return test_connection_endpoint(
            source_config=source_config,
            credentials=credentials,
            connector_id=self.id,
            default_message="Excel connector ready.",
        )

    def discover_objects(
        self,
        source_config: dict[str, Any],
        credentials: dict[str, Any],
        options: dict[str, Any] | None = None,
    ) -> list[DiscoveredObject]:
        opts = options or {}
        path_opt = str(opts.get("path") or "").strip()
        proto = resolve_protocol(source_config, opts)

        # If path points directly to an Excel workbook, discover its sheets
        if path_opt and proto == "local":
            files = resolve_local_file_paths(source_config, opts, default_pattern="*.xlsx")
            if files and files[0].is_file():
                try:
                    import openpyxl  # type: ignore

                    wb = openpyxl.load_workbook(str(files[0]), read_only=True)
                    sheet_names = wb.sheetnames
                    wb.close()
                    return [
                        DiscoveredObject(
                            id=name,
                            name=name,
                            type="sheet",
                            path=str(files[0]),
                            metadata={"sheet_name": name, "file": str(files[0])},
                        )
                        for name in sheet_names
                    ]
                except Exception as exc:
                    raise DiscoveryError(f"Failed to inspect sheets in {files[0]}: {exc}", connector_id=self.id) from exc

        # Otherwise, discover files in the folder or server
        return discover_files_endpoint(
            source_config=source_config,
            credentials=credentials,
            default_pattern="*.xlsx",
            options=options,
            connector_id=self.id,
        )

    def _read_workbook_source(
        self,
        source: str | io.BytesIO,
        table_options: dict[str, Any],
        n_rows: int | None = None,
    ) -> pl.DataFrame:
        sheet = table_options.get("sheet_name") or None
        has_header = bool(table_options.get("has_header", True))
        skip_rows = int(table_options.get("skip_rows") or 0)
        cols_raw = str(table_options.get("columns") or "").strip()
        selected_cols = [c.strip() for c in cols_raw.split(",") if c.strip()] if cols_raw else None

        if skip_rows > 0:
            import openpyxl

            wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
            try:
                ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb.active
                raw_rows = list(ws.iter_rows(values_only=True))
                if skip_rows < len(raw_rows):
                    raw_rows = raw_rows[skip_rows:]
                else:
                    raw_rows = []
                if not raw_rows:
                    df = pl.DataFrame()
                elif has_header:
                    header = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(raw_rows[0])]
                    data_rows = raw_rows[1:]
                    data_rows = [list(r) + [None] * (len(header) - len(r)) for r in data_rows]
                    data_dict = {col: [r[i] for r in data_rows] for i, col in enumerate(header)}
                    df = pl.DataFrame(data_dict, strict=False)
                else:
                    header = [f"col_{i}" for i in range(len(raw_rows[0]))]
                    data_rows = [list(r) + [None] * (len(header) - len(r)) for r in raw_rows]
                    data_dict = {col: [r[i] for r in data_rows] for i, col in enumerate(header)}
                    df = pl.DataFrame(data_dict, strict=False)
            finally:
                wb.close()
        else:
            try:
                df = pl.read_excel(
                    source,
                    sheet_name=sheet,
                    has_header=has_header,
                    engine="calamine",
                )
            except Exception:
                if isinstance(source, io.BytesIO):
                    source.seek(0)
                df = pl.read_excel(
                    source,
                    sheet_name=sheet,
                    has_header=has_header,
                    engine="openpyxl",
                )

        if selected_cols:
            valid_cols = [c for c in selected_cols if c in df.columns]
            if valid_cols:
                df = df.select(valid_cols)

        if n_rows is not None:
            df = df.head(n_rows)

        return df

    def _read_polars_df(
        self,
        source_config: dict[str, Any],
        credentials: dict[str, Any],
        table_options: dict[str, Any],
        n_rows: int | None = None,
    ) -> pl.DataFrame:
        proto = resolve_protocol(source_config, table_options)

        if proto == "local":
            files = resolve_local_file_paths(source_config, table_options, default_pattern="*.xlsx")
            if not files:
                raise SchemaInferenceError("No matching Excel files found for path.", connector_id=self.id)
            if not files[0].exists():
                raise SchemaInferenceError(f"Excel file does not exist: {files[0]}", connector_id=self.id)

            if len(files) == 1:
                try:
                    return self._read_workbook_source(str(files[0]), table_options, n_rows=n_rows)
                except Exception as exc:
                    raise SchemaInferenceError(f"Failed to read Excel workbook: {exc}", connector_id=self.id) from exc
            else:
                dfs: list[pl.DataFrame] = []
                for f in files:
                    try:
                        df_part = self._read_workbook_source(str(f), table_options, n_rows=n_rows)
                        dfs.append(df_part)
                    except Exception as exc:
                        raise SchemaInferenceError(f"Failed reading Excel file {f.name}: {exc}", connector_id=self.id) from exc

                if not dfs:
                    raise SchemaInferenceError("No data could be read from Excel files.", connector_id=self.id)
                combined = pl.concat(dfs, how="vertical_relaxed")
                if n_rows is not None:
                    combined = combined.head(n_rows)
                return combined
        else:
            raw_bytes = fetch_remote_bytes(source_config, credentials, table_options, connector_id=self.id)
            try:
                buf = io.BytesIO(raw_bytes)
                return self._read_workbook_source(buf, table_options, n_rows=n_rows)
            except Exception as exc:
                raise SchemaInferenceError(f"Failed to parse Excel from remote bytes: {exc}", connector_id=self.id) from exc

    def infer_schema(
        self,
        source_config: dict[str, Any],
        credentials: dict[str, Any],
        table_options: dict[str, Any],
    ) -> SchemaResult:
        self.validate_table_options(table_options)
        try:
            df = self._read_polars_df(source_config, credentials, table_options, n_rows=1000)
        except Exception as exc:
            if isinstance(exc, (SchemaInferenceError, PreviewError)):
                raise
            raise SchemaInferenceError(f"Excel schema inference failed: {exc}", connector_id=self.id) from exc

        columns = [
            map_polars_dtype_to_metadata(name, dtype, is_nullable=df[name].null_count() > 0 or len(df) == 0)
            for name, dtype in zip(df.columns, df.dtypes)
        ]
        return SchemaResult(
            columns=columns,
            row_count_estimate=len(df),
            sample_size=len(df),
        )

    def preview(
        self,
        source_config: dict[str, Any],
        credentials: dict[str, Any],
        table_options: dict[str, Any],
        limit: int = 50,
    ) -> PreviewResult:
        self.validate_table_options(table_options)
        try:
            df = self._read_polars_df(source_config, credentials, table_options, n_rows=limit)
        except Exception as exc:
            if isinstance(exc, PreviewError):
                raise
            raise PreviewError(f"Excel preview failed: {exc}", connector_id=self.id) from exc

        rows = format_preview_rows(df)
        return PreviewResult(columns=df.columns, rows=rows, total_rows=len(rows))
